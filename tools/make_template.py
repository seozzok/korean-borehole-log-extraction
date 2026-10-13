# -*- coding: utf-8 -*-
"""Create the Excel export template (assets/__template.xlsx).

Layout expected by ui/utils.write_on_xlsx: one sheet per table, data written from row 6, column B.
Row 1: table title, row 3: Korean column names (as shown in the program), row 4: English column names,
row 5: units.

usage: python tools/make_template.py
"""
import os

import openpyxl
from openpyxl.styles import Font, PatternFill

# sheet: (title, [(Korean name, English name, unit), ...])
SHEETS = {
    'LOCA': ('Location details', [
        ('시추공', 'Borehole', ''), ('Easting', 'Easting', 'm'), ('Northing', 'Northing', 'm'),
        ('최종시추심도', 'Final depth', 'm'), ('시추위치 지반표고', 'Ground elevation', 'EL.m'),
        ('최종굴착심도', 'Final excavation depth', 'm'), ('작업 일시', 'Date', 'yyyy-mm-dd')]),
    'WSTD': ('Groundwater level', [
        ('시추공', 'Borehole', ''), ('지하수위', 'Groundwater level', 'GL.-m')]),
    'SAMP': ('Sample information', [
        ('시추공', 'Borehole', ''), ('시료 상단심도', 'Sample top', 'm'), ('시료 하단심도', 'Sample base', 'm'),
        ('시료번호', 'Sample ID', ''), ('시료타입', 'Sample type', ''), ('시료채취 일시', 'Sampling date', ''),
        ('SPT 값', 'SPT blow count at sampling', ''), ('시료 지름', 'Sample diameter', 'mm'),
        ('지하수위', 'Groundwater level', 'GL.-m'), ('시료 설명', 'Sample description', ''),
        ('비고', 'Remarks', '')]),
    'CORE': ('Core information', [
        ('시추공', 'Borehole', ''), ('코어시료 상단심도', 'Core top', 'm'), ('코어시료 하단심도', 'Core base', 'm'),
        ('TCR값', 'TCR', '%'), ('RQD값', 'RQD', '%'), ('코어시료 지름', 'Core diameter', 'mm')]),
    'CDIA': ('Casing diameter', [
        ('시추공', 'Borehole', ''), ('케이싱 심도', 'Casing depth', 'm'), ('케이싱 지름', 'Casing diameter', 'mm')]),
    'GEOL': ('Geological layers', [
        ('시추공', 'Borehole', ''), ('지층분류', 'Layer', ''), ('지층상단심도', 'Layer top', 'm'),
        ('지층하단심도', 'Layer base', 'm'), ('지층종류', 'Soil/rock type (USCS)', ''),
        ('지층순서', 'Layer order', ''), ('비고', 'Layer description', '')]),
    'DISC': ('Discontinuity data', [
        ('시추공', 'Borehole', ''), ('상단심도', 'Top', 'm'), ('하단심도', 'Base', 'm'),
        ('주절리군 번호', 'Joint set', ''), ('불연속면 번호', 'Discontinuity no.', ''),
        ('불연속면 종류', 'Discontinuity type', ''), ('불연속면 경사', 'Dip', 'deg'),
        ('불연속면 경사방향', 'Dip direction', 'deg'), ('소규모/중규모 거칠기', 'Roughness', ''),
        ('충전물질', 'Infilling', ''), ('Core Loss 유무', 'Core loss', ''), ('파쇄 유무', 'Fractured zone', ''),
        ('암맥 (또는 맥암) 유무', 'Dike', ''), ('단층 유무', 'Fault', '')]),
    'WETH': ('Weathering', [
        ('시추공', 'Borehole', ''), ('상단심도', 'Top', 'm'), ('하단심도', 'Base', 'm'),
        ('풍화등급', 'Weathering grade', '')]),
    'FRAC': ('Fracture spacing', [
        ('시추공', 'Borehole', ''), ('상단심도', 'Top', 'm'), ('하단심도', 'Base', 'm'),
        ('주절리군 번호', 'Joint set', ''), ('절리 최대간격', 'Maximum spacing', 'mm'),
        ('절리 최소간격', 'Minimum spacing', 'mm'), ('절리 평균간격', 'Mean spacing', 'mm'),
        ('FI 값 (m당 불연속면수)', 'Fracture index', '1/m'), ('절리간격 분류', 'Spacing class', '')]),
    'ISPT': ('Standard penetration test', [
        ('시추공', 'Borehole', ''), ('시험 깊이', 'Test depth', 'm'), ('타수', 'Blow count', ''),
        ('관입깊이', 'Penetration', 'mm'), ('N 값', 'N value', '')]),
}


def main(out_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    fill = PatternFill('solid', fgColor='DDEBF7')
    for name, (title, cols) in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.cell(1, 2, title).font = Font(bold=True, size=12)
        for j, (ko, en, unit) in enumerate(cols, start=2):
            for r, v in ((3, ko), (4, en), (5, unit)):
                c = ws.cell(r, j, v)
                c.fill = fill
                if r < 5:
                    c.font = bold
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = max(12, len(en) + 2)
        ws.cell(5, 1, '<UNITS>')
        ws.freeze_panes = 'B6'
    wb.save(out_path)
    print('saved', out_path)


if __name__ == '__main__':
    here = os.path.dirname(os.path.abspath(__file__))
    main(os.path.join(here, '..', 'assets', '__template.xlsx'))
